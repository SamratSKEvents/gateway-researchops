import json
from pathlib import Path
from pptx import Presentation
from backend.modules.deck import build

RUNS = Path(__file__).resolve().parents[3] / "data" / "runs"


def test_deck_builds_from_a_minimal_run():
    snap = {"query": "q", "claims": {}, "sources": {}, "result": {"report": {"verdict": {"label": "favourable", "p": 0.7, "low": 0.6, "high": 0.8},
            "hypotheses": [{"id": "h1", "text": "demand", "p": 0.7}], "must_do": [{"text": "pilot"}]}}}
    prs = Presentation(build(snap))
    assert len(prs.slides) == 10
    assert any("pilot" in sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
